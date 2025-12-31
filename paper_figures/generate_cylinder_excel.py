"""
Generate cylinder experiment results Excel file.
One sheet per backbone model (FNO, CNO, DeepONet, DPOT-S).
Each row = one training method (Pretrain, Vanilla FT, L2-SP, EWC, NSFT).
Each column = one evaluation metric (best-RMSE checkpoint).
"""
import re
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "./results"

# ─────────────────────────────────────────────────────────────────────────────
# Log parsers
# ─────────────────────────────────────────────────────────────────────────────

def parse_new_format(lines):
    """
    Parser for train_nullspace.py / newer train.py logs.
    Pattern:
        Iteration X, train loss: Y[, protection_beta: Z]
        (timestamp) - INFO - Validation results:
        normalized mse loss: V1, rmse: V2, ...
    """
    records = []
    i = 0
    while i < len(lines):
        # Find "Iteration X" line
        m = re.search(r"Iteration\s+(\d+),\s*train loss:\s*([\d.]+)", lines[i])
        if m:
            iteration = int(m.group(1))
            # Look ahead for "Validation results:"
            j = i + 1
            while j < min(i + 5, len(lines)):
                if "Validation results:" in lines[j]:
                    # Metrics are on the next non-empty line
                    k = j + 1
                    while k < min(j + 5, len(lines)):
                        metric_line = lines[k].strip()
                        if metric_line and not metric_line.startswith("20"):
                            rec = _parse_new_metric_line(iteration, metric_line)
                            if rec:
                                records.append(rec)
                            break
                        k += 1
                    break
                j += 1
        i += 1
    return records


def _parse_new_metric_line(iteration, line):
    """Parse the comma-separated metrics line from new format."""
    kv = {}
    for part in line.split(","):
        part = part.strip()
        # "normalized mse loss: 0.123" or "rmse: 0.123"
        m = re.match(r"^(.+?):\s*([\d.]+(?:e[+-]?\d+)?)$", part)
        if m:
            key = m.group(1).strip().lower()
            val = float(m.group(2))
            kv[key] = val
    if "rmse" not in kv:
        return None
    return {
        "iteration":     iteration,
        "norm_mse":      kv.get("normalized mse loss", float("nan")),
        "rmse":          kv.get("rmse", float("nan")),
        "mae":           kv.get("mae", float("nan")),
        "rel_l2":        kv.get("rel l2 error", float("nan")),
        "r2":            kv.get("r2", float("nan")),
        "ke_err":        kv.get("ke error", float("nan")),
        "f_err":         kv.get("f error", float("nan")),
        "low_f":         kv.get("low f error", float("nan")),
        "mid_f":         kv.get("mid f error", float("nan")),
        "high_f":        kv.get("high f error", float("nan")),
    }


def parse_old_format(lines):
    """
    Parser for old train.py / train_gpus.py logs.
    Pattern:
        Iteration X, train loss: Y
        Validation | normalized_mse=V1  rmse=V2  mae=V3  rel_l2=V4  r2=V5
          ke_err=V6  f_err=V7  low_f=V8  mid_f=V9  high_f=V10
    """
    records = []
    i = 0
    while i < len(lines):
        m = re.search(r"Iteration\s+(\d+),\s*train loss:\s*([\d.]+)", lines[i])
        if m:
            iteration = int(m.group(1))
            j = i + 1
            while j < min(i + 5, len(lines)):
                if lines[j].startswith("Validation |"):
                    # Collect continuation lines (lines starting with spaces)
                    val_text = lines[j]
                    k = j + 1
                    while k < min(j + 4, len(lines)) and (lines[k].startswith("  ") or lines[k].startswith("\t")):
                        val_text += " " + lines[k].strip()
                        k += 1
                    rec = _parse_old_metric_line(iteration, val_text)
                    if rec:
                        records.append(rec)
                    break
                j += 1
        i += 1
    return records


def _parse_old_metric_line(iteration, text):
    """Parse space-separated key=value metrics from old format."""
    kv = {}
    for m in re.finditer(r"([\w_]+)=([\d.]+(?:e[+-]?\d+)?)", text):
        kv[m.group(1).lower()] = float(m.group(2))
    if "rmse" not in kv:
        return None
    return {
        "iteration": iteration,
        "norm_mse":  kv.get("normalized_mse", float("nan")),
        "rmse":      kv.get("rmse", float("nan")),
        "mae":       kv.get("mae", float("nan")),
        "rel_l2":    kv.get("rel_l2", float("nan")),
        "r2":        kv.get("r2", float("nan")),
        "ke_err":    kv.get("ke_err", float("nan")),
        "f_err":     kv.get("f_err", float("nan")),
        "low_f":     kv.get("low_f", float("nan")),
        "mid_f":     kv.get("mid_f", float("nan")),
        "high_f":    kv.get("high_f", float("nan")),
    }


def load_log(path):
    with open(path, "r", errors="ignore") as f:
        return f.readlines()


def detect_format(lines):
    for line in lines[:50]:
        if "Validation |" in line:
            return "old"
        if "Validation results:" in line:
            return "new"
    return "new"


def parse_log(path):
    lines = load_log(path)
    fmt = detect_format(lines)
    if fmt == "old":
        return parse_old_format(lines)
    else:
        return parse_new_format(lines)


def best_record(records):
    """Return the record with the lowest RMSE."""
    if not records:
        return None
    return min(records, key=lambda r: r["rmse"])


def best_from_logs(log_paths):
    """Combine records from multiple logs and return the best."""
    all_records = []
    for path in log_paths:
        if os.path.exists(path):
            all_records.extend(parse_log(path))
    return best_record(all_records)


# ─────────────────────────────────────────────────────────────────────────────
# Experiment definitions
# ─────────────────────────────────────────────────────────────────────────────

EXPERIMENTS = {
    "FNO": [
        {
            "method": "Pretrain (Num.)",
            "logs": [f"{BASE}/fno/fno_cylinder_pretrained/2026-03-09_17-33-29/training.log"],
        },
        {
            "method": "Vanilla FT",
            "logs": [f"{BASE}/fno/fno_cylinder_dft/2026-03-10_03-40-24/training.log"],
        },
        {
            "method": "L2-SP",
            "logs": [f"{BASE}/fno/fno_cylinder_dft/2026-03-10_03-40-54/training.log"],
        },
        {
            "method": "EWC",
            "logs": [f"{BASE}/fno/fno_cylinder_ewcft/2026-03-12_20-09-12/training.log"],
        },
        {
            "method": "NSFT (Ours)",
            "logs": [
                f"{BASE}/fno/fno_cylinder_nsft/2026-03-07_11-01-08/training.log",
                f"{BASE}/fno/fno_cylinder_nsft/2026-03-08_10-15-26/training.log",
                f"{BASE}/fno/fno_cylinder_nsft/2026-03-10_08-42-11/training.log",
                f"{BASE}/fno/fno_cylinder_nsft/2026-03-12_22-38-44/training.log",
                f"{BASE}/fno/fno_cylinder_nsft/2026-03-13_02-09-31/training.log",
            ],
            "highlight": True,
        },
    ],
    "CNO": [
        {
            "method": "Pretrain (Num.)",
            "logs": [f"{BASE}/cno/cno_cylinder_pretrained/2026-03-11_09-49-04/training.log"],
        },
        {
            "method": "NSFT (Ours)",
            "logs": [f"{BASE}/cno/cno_cylinder_nsft/2026-03-13_00-37-38/training.log"],
            "highlight": True,
        },
    ],
    "DeepONet": [
        {
            "method": "Pretrain (Num.)",
            "logs": [f"{BASE}/deeponet/deeponet_cylinder_pretrained/2026-03-12_11-16-03/training.log"],
        },
        {
            "method": "NSFT (Ours)",
            "logs": [f"{BASE}/deeponet/deeponet_cylinder_nsft/2026-03-13_00-38-45/training.log"],
            "highlight": True,
        },
    ],
    "DPOT-S": [
        {
            "method": "Pretrain (Num.)",
            "logs": [f"{BASE}/dpot/dpot_s_cylinder_pretrained/2026-03-12_01-49-29/training.log"],
        },
        {
            "method": "NSFT (Ours, partial)",
            "logs": [f"{BASE}/dpot/dpot_s_cylinder_nsft/2026-03-13_09-41-17/training.log"],
            "highlight": True,
        },
    ],
}

COLUMNS = [
    ("Method",          None),
    ("Best Iter",       None),
    ("RMSE ↓",          "rmse"),
    ("MAE ↓",           "mae"),
    ("Rel L2 ↓",        "rel_l2"),
    ("R² ↑",            "r2"),
    ("KE Error ↓",      "ke_err"),
    ("F Error ↓",       "f_err"),
    ("Low-F Error ↓",   "low_f"),
    ("Mid-F Error ↓",   "mid_f"),
    ("High-F Error ↓",  "high_f"),
    ("Norm MSE ↓",      "norm_mse"),
]


# ─────────────────────────────────────────────────────────────────────────────
# Excel styling
# ─────────────────────────────────────────────────────────────────────────────

HEADER_FILL  = PatternFill("solid", fgColor="2E4057")   # dark blue
OURS_FILL    = PatternFill("solid", fgColor="D6EAF8")   # light blue
ALT_FILL     = PatternFill("solid", fgColor="F8F9FA")   # light grey
BEST_FONT    = Font(bold=True, color="C0392B")          # red bold for best per column

THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")


def style_header(cell):
    cell.font      = Font(bold=True, color="FFFFFF", size=10)
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER
    cell.border    = BORDER


def style_data(cell, highlight=False, alt=False):
    if highlight:
        cell.fill = OURS_FILL
    elif alt:
        cell.fill = ALT_FILL
    cell.alignment = CENTER
    cell.border    = BORDER


def write_sheet(wb, sheet_name, experiments):
    ws = wb.create_sheet(title=sheet_name)

    # Write title
    ws.merge_cells(f"A1:{get_column_letter(len(COLUMNS))}1")
    title_cell = ws["A1"]
    title_cell.value = f"Cylinder Dataset — {sheet_name} Results"
    title_cell.font  = Font(bold=True, size=13, color="2E4057")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Write column headers (row 2)
    for col_idx, (col_name, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=col_name)
        style_header(cell)
    ws.row_dimensions[2].height = 18

    # Parse all records
    rows_data = []
    for exp in experiments:
        rec = best_from_logs(exp["logs"])
        rows_data.append((exp, rec))

    # Find per-column best values for highlighting
    col_best = {}
    for col_name, key in COLUMNS[2:]:   # skip Method, Best Iter
        vals = [rec[key] for _, rec in rows_data if rec and key in rec]
        if vals:
            if "↑" in col_name:
                col_best[key] = max(vals)
            else:
                col_best[key] = min(vals)

    # Write data rows
    for row_idx, (exp, rec) in enumerate(rows_data, start=3):
        highlight = exp.get("highlight", False)
        alt       = (row_idx % 2 == 0)

        # Method name
        cell = ws.cell(row=row_idx, column=1, value=exp["method"])
        cell.alignment = LEFT
        if highlight:
            cell.fill = OURS_FILL
            cell.font = Font(bold=True)
        elif alt:
            cell.fill = ALT_FILL
        cell.border = BORDER

        if rec is None:
            ws.cell(row=row_idx, column=2, value="—").border = BORDER
            for col_idx in range(3, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx, value="—").border = BORDER
            continue

        # Best Iter
        cell = ws.cell(row=row_idx, column=2, value=rec["iteration"])
        style_data(cell, highlight, alt)

        # Metric columns
        for col_idx, (col_name, key) in enumerate(COLUMNS[2:], start=3):
            val = rec.get(key, float("nan"))
            if val != val:   # nan check
                text = "—"
            else:
                text = f"{val:.5f}"
            cell = ws.cell(row=row_idx, column=col_idx, value=text)
            style_data(cell, highlight, alt)
            # Bold+red if this is the column's best value
            if key in col_best and abs(val - col_best[key]) < 1e-9:
                cell.font = BEST_FONT if not highlight else Font(bold=True, color="C0392B")

        ws.row_dimensions[row_idx].height = 16

    # Column widths
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    for col_idx in range(3, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    # Freeze header rows
    ws.freeze_panes = "A3"


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default empty sheet

    for sheet_name, experiments in EXPERIMENTS.items():
        print(f"Generating sheet: {sheet_name}")
        write_sheet(wb, sheet_name, experiments)

        # Print summary to console
        for exp in experiments:
            rec = best_from_logs(exp["logs"])
            if rec:
                print(f"  {exp['method']:25s} iter={rec['iteration']:5d}  "
                      f"rmse={rec['rmse']:.5f}  mae={rec['mae']:.5f}  "
                      f"r2={rec['r2']:.5f}")
            else:
                print(f"  {exp['method']:25s} — no data")

    out_path = "./cylinder_results.xlsx"
    wb.save(out_path)
    print(f"\nSaved: {out_path}")


if __name__ == "__main__":
    main()
