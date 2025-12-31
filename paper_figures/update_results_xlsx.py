#!/usr/bin/env python3
"""Collect all cylinder training results and write to cylinder_results.xlsx
Structure: Summary sheet (first) + one sheet per model with all metrics.
Best value per metric is bold+red; running experiments are yellow-highlighted.
"""
import re, glob, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

RESULTS   = "./results"
XLSX_PATH = "./cylinder_results.xlsx"

MODEL_ORDER  = ['FNO', 'CNO', 'DPOT-S', 'DeepONet']
METHOD_ORDER = ['Pretrained (Sim-only)', 'DFT', 'L2-FT', 'EWC-FT', 'NSFT']

# All metrics: (col_key, display_name, lower_is_better)
METRICS = [
    ('rmse',          'RMSE',          True),
    ('normalized_mse','Norm MSE',      True),
    ('mae',           'MAE',           True),
    ('rel_l2',        'Rel L2',        True),
    ('r2',            'R2',            False),
    ('ke_error',      'KE Error',      True),
    ('f_error',       'F Error',       True),
    ('low_f',         'Low-F Err',     True),
    ('mid_f',         'Mid-F Err',     True),
    ('high_f',        'High-F Err',    True),
    ('freq_error',    'Freq Error',    True),
]

# ── Parsing helpers ───────────────────────────────────────
def parse_best_metrics(log_path):
    with open(log_path) as f:
        content = f.read()
    # Find explicitly logged best RMSE
    m = re.search(r'Best val RMSE[:\s]+([0-9.]+)', content)
    best_rmse_explicit = float(m.group(1)) if m else None

    # Collect all validation blocks
    blocks = []
    lines = content.splitlines()
    for i, line in enumerate(lines):
        metrics = {}
        # Format A (finetune): "normalized mse loss: X, rmse: X, ..."
        if re.search(r'normalized mse loss:', line):
            combined = line
            if i+1 < len(lines):
                nxt = lines[i+1]
                if not re.match(r'\d{4}-\d{2}-\d{2}', nxt.strip()) and 'INFO' not in nxt:
                    combined += ' ' + nxt
            for pat, key in [
                (r'normalized mse loss[:\s]+([0-9.]+)',  'normalized_mse'),
                (r'\brmse[:\s]+([0-9.]+)',               'rmse'),
                (r'\bmae[:\s]+([0-9.]+)',                'mae'),
                (r'rel l2 error[:\s]+([0-9.]+)',         'rel_l2'),
                (r'\br2[:\s]+([0-9.]+)',                 'r2'),
                (r'ke error[:\s]+([0-9.]+)',             'ke_error'),
                (r'\bf error[:\s]+([0-9.]+)',            'f_error'),
                (r'low f error[:\s]+([0-9.]+)',          'low_f'),
                (r'mid f error[:\s]+([0-9.]+)',          'mid_f'),
                (r'high f error[:\s]+([0-9.]+)',         'high_f'),
                (r'freq error[:\s]+([0-9.]+)',           'freq_error'),
            ]:
                mm = re.search(pat, combined)
                if mm: metrics[key] = float(mm.group(1))
        # Format B (pretrained): "Validation | normalized_mse=X  rmse=X ..."
        elif 'Validation |' in line and 'normalized_mse=' in line:
            combined = line
            for j in (1, 2):
                if i+j < len(lines): combined += ' ' + lines[i+j]
            for pat, key in [
                (r'normalized_mse=([0-9.]+)',  'normalized_mse'),
                (r'\brmse=([0-9.]+)',           'rmse'),
                (r'\bmae=([0-9.]+)',            'mae'),
                (r'rel_l2=([0-9.]+)',           'rel_l2'),
                (r'\br2=([0-9.]+)',             'r2'),
                (r'ke_err=([0-9.]+)',           'ke_error'),
                (r'f_err=([0-9.]+)',            'f_error'),
                (r'low_f=([0-9.]+)',            'low_f'),
                (r'mid_f=([0-9.]+)',            'mid_f'),
                (r'high_f=([0-9.]+)',           'high_f'),
                (r'freq_err=([0-9.]+)',         'freq_error'),
            ]:
                mm = re.search(pat, combined)
                if mm: metrics[key] = float(mm.group(1))
        if metrics:
            blocks.append(metrics)

    if not blocks:
        return {}
    target = best_rmse_explicit
    if target is None:
        target = min(b.get('rmse', 1e9) for b in blocks)
    best_block = min(blocks, key=lambda b: abs(b.get('rmse', 1e9) - target))
    return best_block


def get_best_iter(log_path):
    with open(log_path) as f: content = f.read()
    m = re.search(r'[Bb]est iteration[:\s=]+(\d+)', content)
    return int(m.group(1)) if m else None

def get_final_iter(log_path):
    last = None
    with open(log_path) as f:
        for line in f:
            mm = re.match(r'Iteration (\d+)', line)
            if mm: last = int(mm.group(1))
    return last

def get_reg_lambda(log_path):
    with open(log_path) as f: first = f.read(3000)
    m = re.search(r'reg_lambda=([0-9.e+-]+)', first)
    return m.group(1) if m else None

def is_complete(log_path):
    with open(log_path) as f: content = f.read()
    return any(k in content for k in ['Training complete','fine-tuning complete','训练完成'])

# ── Collect all experiments ───────────────────────────────
rows = []
for log_path in sorted(glob.glob(f"{RESULTS}/*/**/training.log", recursive=True)):
    parts = log_path.replace(RESULTS+"/","").split("/")
    model_dir = parts[0]
    exp_name  = parts[1]
    model = {'fno':'FNO','cno':'CNO','dpot':'DPOT-S','deeponet':'DeepONet'}.get(model_dir, model_dir.upper())
    if   'nsft'       in exp_name: method, lam = 'NSFT',                  '-'
    elif 'pretrained' in exp_name: method, lam = 'Pretrained (Sim-only)', '-'
    elif 'ewcft'      in exp_name: method, lam = 'EWC-FT',  get_reg_lambda(log_path) or '?'
    elif 'l2ft'       in exp_name: method, lam = 'L2-FT',   get_reg_lambda(log_path) or '?'
    elif 'dft'        in exp_name: method, lam = 'DFT',                   '-'
    else:                          method, lam = exp_name,                 '-'
    metrics = parse_best_metrics(log_path)
    row = {
        'Model':      model,
        'Method':     method,
        'λ':          lam,
        'Best Iter':  get_best_iter(log_path),
        'Final Iter': get_final_iter(log_path),
        'Status':     'Done' if is_complete(log_path) else 'Running',
    }
    row.update(metrics)
    rows.append(row)

rows.sort(key=lambda r:(
    MODEL_ORDER.index(r['Model'])   if r['Model']  in MODEL_ORDER  else 99,
    METHOD_ORDER.index(r['Method']) if r['Method'] in METHOD_ORDER else 99,
    str(r['λ'])))

# ── Style constants ───────────────────────────────────────
thin   = Side(style='thin',   color='BFBFBF')
medium = Side(style='medium', color='888888')
bdr    = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr    = Alignment(horizontal='center', vertical='center', wrap_text=True)
HDR    = "1F4E79"
BASE   = "D9E1F2"   # light blue  – Pretrained
NSFT_C = "E2EFDA"   # light green – NSFT
RUN    = "FFF2CC"   # yellow      – Running
WHT    = "FFFFFF"
BEST_FONT = Font(bold=True, color="C00000")

def row_bg(r):
    if r['Status'] == 'Running':    return RUN
    if 'NSFT'       in r['Method']: return NSFT_C
    if 'Pretrained' in r['Method']: return BASE
    return WHT

def fmt(v):
    if v is None: return '-'
    if isinstance(v, float): return round(v, 6)
    return v

def style_header(ws, row_idx, ncols, height=28):
    ws.row_dimensions[row_idx].height = height
    for col in range(1, ncols+1):
        c = ws.cell(row_idx, col)
        c.font      = Font(bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", fgColor=HDR)
        c.alignment = ctr
        c.border    = bdr

def apply_best_highlight(ws, data_row_start, data_row_end, col_idx, lower_better):
    vals = []
    for ri in range(data_row_start, data_row_end+1):
        v = ws.cell(ri, col_idx).value
        if isinstance(v, float): vals.append((v, ri))
    if not vals: return
    best_ri = (min if lower_better else max)(vals, key=lambda x: x[0])[1]
    ws.cell(best_ri, col_idx).font = BEST_FONT

# ══════════════════════════════════════════════════════════
# Build workbook
now_str = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
wb = openpyxl.Workbook()
wb.remove(wb.active)   # remove default sheet

# ── SUMMARY SHEET (first) ─────────────────────────────────
ws_sum = wb.create_sheet("Summary")
sum_hdr = ['Model', 'Method', 'λ', 'Status',
           'RMSE', 'Norm MSE', 'MAE', 'Rel L2', 'R2',
           'KE Err', 'F Err', 'Freq Err', 'Best Iter', 'Updated']
ws_sum.append(sum_hdr)
style_header(ws_sum, 1, len(sum_hdr))

SUM_METRIC_KEYS = ['rmse','normalized_mse','mae','rel_l2','r2','ke_error','f_error','freq_error']
SUM_METRIC_COL0 = 5  # 1-based column of first metric

for r in rows:
    vals = [r['Model'], r['Method'], r['λ'], r['Status']]
    for k in SUM_METRIC_KEYS: vals.append(fmt(r.get(k)))
    vals += [r['Best Iter'], now_str]
    ws_sum.append(vals)
    ri = ws_sum.max_row
    bg = row_bg(r)
    for col in range(1, len(sum_hdr)+1):
        c = ws_sum.cell(ri, col)
        c.fill      = PatternFill("solid", fgColor=bg)
        c.alignment = ctr
        c.border    = bdr
        if isinstance(c.value, float): c.number_format = '0.000000'

# Best highlight per model group
for model in MODEL_ORDER:
    model_row_idxs = [i+2 for i, r in enumerate(rows) if r['Model'] == model]
    if len(model_row_idxs) < 2: continue
    for mi, k in enumerate(SUM_METRIC_KEYS):
        col = SUM_METRIC_COL0 + mi
        lb = next((lb for ck,_,lb in METRICS if ck==k), True)
        apply_best_highlight(ws_sum, model_row_idxs[0], model_row_idxs[-1], col, lb)

ws_sum.column_dimensions['A'].width = 11
ws_sum.column_dimensions['B'].width = 24
for i in range(3, 14): ws_sum.column_dimensions[get_column_letter(i)].width = 11
ws_sum.column_dimensions[get_column_letter(14)].width = 17
ws_sum.freeze_panes = "A2"

# ── PER-MODEL SHEETS ──────────────────────────────────────
detail_hdr = (['Method', 'λ', 'Status', 'Best Iter', 'Final Iter'] +
               [name for _, name, _ in METRICS])
METRIC_COL0 = 6  # 1-based column of first metric in per-model sheet

for model in MODEL_ORDER:
    model_rows = [r for r in rows if r['Model'] == model]
    if not model_rows: continue

    ws_m = wb.create_sheet(model)

    # Title row
    ncols = len(detail_hdr)
    ws_m.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    tc = ws_m.cell(1, 1)
    tc.value     = f"{model} — Cylinder Fine-tuning Results   (updated {now_str})"
    tc.font      = Font(bold=True, color="FFFFFF", size=12)
    tc.fill      = PatternFill("solid", fgColor=HDR)
    tc.alignment = Alignment(horizontal='center', vertical='center')
    ws_m.row_dimensions[1].height = 24

    # Header row (row 2)
    ws_m.append(detail_hdr)
    style_header(ws_m, 2, ncols)

    # Data rows start at row 3
    for r in model_rows:
        vals = [r['Method'], r['λ'], r['Status'], r['Best Iter'], r['Final Iter']]
        for k, _, _ in METRICS: vals.append(fmt(r.get(k)))
        ws_m.append(vals)
        ri = ws_m.max_row
        bg = row_bg(r)
        for col in range(1, ncols+1):
            c = ws_m.cell(ri, col)
            c.fill      = PatternFill("solid", fgColor=bg)
            c.alignment = ctr
            c.border    = bdr
            if isinstance(c.value, float): c.number_format = '0.000000'

    # Best highlight per metric column
    data_start, data_end = 3, ws_m.max_row
    for mi, (k, _, lb) in enumerate(METRICS):
        apply_best_highlight(ws_m, data_start, data_end, METRIC_COL0 + mi, lb)

    # Column widths
    fixed_w = [26, 8, 9, 10, 10]
    metric_w = [11] * len(METRICS)
    for i, w in enumerate(fixed_w + metric_w, 1):
        ws_m.column_dimensions[get_column_letter(i)].width = w
    ws_m.freeze_panes = "A3"

wb.save(XLSX_PATH)

sheet_names = ["Summary"] + [m for m in MODEL_ORDER if any(r['Model']==m for r in rows)]
print(f"Saved: {XLSX_PATH}")
print(f"Sheets: {sheet_names}")
print(f"Total rows: {len(rows)}")
